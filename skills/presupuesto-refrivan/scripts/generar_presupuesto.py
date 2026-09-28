"""Genera un presupuesto REFRIVAN (xlsx + pdf) a partir de un JSON.

Uso:
    python generar_presupuesto.py datos.json CARPETA_SALIDA [--png]

El JSON:
{
  "numero": 129,
  "fecha": "28/09/2026",                      # opcional, por defecto hoy
  "cliente": {"nombre": "Marcelo", "contacto": "", "telefono": "",
              "direccion": "", "localidad": "", "provincia": ""},
  "items": [
    {"descripcion": "Heladera vertical de 5 puertas", "modelo": "Exhibidora",
     "cantidad": 4, "precio_unitario": 5000000},
    {"descripcion": "Unidad condensadora sin compresor", "modelo": "1 HP",
     "cantidad": 1, "precio_unitario": 1850000,
     "incluye": ["Condensador con forzador", "Armado y conexión con el compresor del cliente"]}
  ],
  "precio_global": null,                      # un importe único para todo el trabajo
  "iva": 21,                                  # 0, 21 o 10.5
  "cotizacion_usd": 1535,
  "condicion_pago": "100% contado",
  "observaciones": ""
}

Arma una o más hojas A4: si los ítems no entran, corta entre ítems (nunca parte un
producto de sus componentes), la hoja siguiente lleva un encabezado de continuación
y la tabla repite su título, y los totales nunca quedan solos en una hoja.
"""
import datetime
import json
import math
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

AQUI = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(AQUI, "..", "assets", "logo.png")

SLATE, RED, ORANGE, TEAL = "283137", "BF222D", "EB9E06", "19988A"
GREY, LINE, TINT = "5B6670", "D5DBE0", "F1F4F6"
MONEY = '[$-2C0A]"$ "#,##0'

EMPRESA_1 = "Marco Avellaneda 960, Remedios de Escalada, Buenos Aires"
EMPRESA_2 = "Tel +54 011 3840-0340  ·  refrivansitioweb@gmail.com  ·  www.refrivan.com.ar"

COLS = {"A": 6, "B": 32, "C": 8, "D": 8, "E": 12, "F": 7, "G": 14, "H": 18}
PAGE_PT = 752  # alto imprimible de A4 (márgenes 0,45" arriba y 0,8" abajo), en puntos
DESC_CHARS = 54  # caracteres por línea en Descripción (Arial 10)
OBS_CHARS = 115

# alturas de fila en puntos
H_ITEM, H_ITEM_LINE, H_SUB, H_EMPTY = 24, 13, 15, 24


def font(size=10, bold=False, color=SLATE, italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", fgColor=c)


HAIR = Side(style="thin", color=LINE)


def money(v):
    return "$ " + f"{v:,.0f}".replace(",", ".")


def lineas(texto, ancho):
    return max(1, math.ceil(len(texto) / ancho)) if texto else 1


def alto_item(it):
    h = H_ITEM + H_ITEM_LINE * (lineas(it["descripcion"], DESC_CHARS) - 1)
    if it.get("incluye"):
        h += H_SUB * (len(it["incluye"]) + 1) + 4
    return h


class Hoja:
    def __init__(self, datos):
        self.d = datos
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Presupuesto"
        self.ws.sheet_view.showGridLines = False
        self.wb._named_styles["Normal"].font = Font(name="Arial", size=10)
        for col, w in COLS.items():
            self.ws.column_dimensions[col].width = w
        self.r = 1

    # -- utilidades
    def fila(self, alto):
        self.ws.row_dimensions[self.r].height = alto
        self.r += 1
        return self.r - 1

    def celda(self, ref, valor=None, f=None, al=None, fmt=None, relleno=None):
        c = self.ws[ref]
        if valor is not None:
            c.value = valor
        if f:
            c.font = f
        if al:
            c.alignment = al
        if fmt:
            c.number_format = fmt
        if relleno:
            c.fill = fill(relleno)
        return c

    def unir(self, rango):
        self.ws.merge_cells(rango)

    def pintar(self, r, cols, color):
        for col in cols:
            self.ws[f"{col}{r}"].fill = fill(color)

    def franja(self, alto=4):
        r = self.fila(alto)
        self.pintar(r, "ABC", RED)
        self.pintar(r, "DEF", ORANGE)
        self.pintar(r, "GH", TEAL)

    def logo(self, fila, alto_px):
        img = XLImage(LOGO)
        ratio = img.width / img.height
        img.height = alto_px
        img.width = round(alto_px * ratio)
        self.ws.add_image(img, f"A{fila}")

    # -- bloques
    def encabezado(self):
        d = self.d
        self.fila(4)
        r1 = self.fila(20)
        r2 = self.fila(30)
        r3 = self.fila(18)
        self.logo(r1, 66)
        self.unir(f"F{r1}:H{r1}")
        self.celda(f"F{r1}", "P R E S U P U E S T O", font(9, True, GREY), Alignment(horizontal="right", vertical="bottom"))
        self.unir(f"F{r2}:H{r2}")
        self.celda(f"F{r2}", f"N° {d['numero']}", font(22, True), Alignment(horizontal="right", vertical="center"))
        self.unir(f"F{r3}:H{r3}")
        self.celda(f"F{r3}", f"Fecha  {d['fecha']}", font(9, False, GREY), Alignment(horizontal="right", vertical="center"))
        r4 = self.fila(16)
        self.unir(f"A{r4}:H{r4}")
        self.celda(f"A{r4}", f"{EMPRESA_1}  ·  {EMPRESA_2}", font(7.5, False, GREY),
                   Alignment(horizontal="left", vertical="bottom"))
        self.fila(6)
        self.franja(5)
        self.fila(14)

    def encabezado_continuacion(self, hoja, total):
        d = self.d
        self.fila(4)
        r1 = self.fila(17)
        r2 = self.fila(17)
        self.logo(r1, 40)
        self.unir(f"E{r1}:H{r1}")
        self.celda(f"E{r1}", f"PRESUPUESTO  N° {d['numero']}", font(10, True), Alignment(horizontal="right", vertical="center"))
        self.unir(f"E{r2}:H{r2}")
        self.celda(f"E{r2}", f"Continuación  ·  hoja {hoja} de {total}  ·  {d['fecha']}", font(8, False, GREY),
                   Alignment(horizontal="right", vertical="center"))
        self.fila(8)
        self.franja(4)
        self.fila(12)

    def cliente(self):
        c = self.d["cliente"]
        grupos = [("A", "D"), ("E", "F"), ("G", "H")]
        filas = [(("CLIENTE", c.get("nombre")), ("CONTACTO", c.get("contacto")), ("TELÉFONO", c.get("telefono"))),
                 (("DIRECCIÓN", c.get("direccion")), ("LOCALIDAD", c.get("localidad")), ("PROVINCIA", c.get("provincia")))]
        r = self.fila(6)
        self.pintar(r, "ABCDEFGH", TINT)
        for i, fila in enumerate(filas):
            rl = self.fila(13)
            rv = self.fila(20)
            for (a, b), (lab, val) in zip(grupos, fila):
                for rr in (rl, rv):
                    self.pintar(rr, "ABCDEFGH", TINT)
                    self.unir(f"{a}{rr}:{b}{rr}")
                self.celda(f"{a}{rl}", lab, font(7, True, GREY), Alignment(horizontal="left", vertical="bottom", indent=1))
                grande = i == 0 and a == "A"
                self.celda(f"{a}{rv}", val or "", font(11 if grande else 10, grande),
                           Alignment(horizontal="left", vertical="center", indent=1))
            if i == 0:
                r = self.fila(4)
                self.pintar(r, "ABCDEFGH", TINT)
        r = self.fila(6)
        self.pintar(r, "ABCDEFGH", TINT)
        self.fila(14)

    def titulo_tabla(self):
        r = self.fila(22)
        self.unir(f"B{r}:D{r}")
        self.pintar(r, "ABCDEFGH", SLATE)
        for col, txt, h in (("A", "N°", "center"), ("B", "DESCRIPCIÓN", "left"), ("E", "MODELO", "left"),
                            ("F", "CANT.", "center"), ("G", "PRECIO UNIT.", "right"), ("H", "SUBTOTAL", "right")):
            self.celda(f"{col}{r}", txt, font(8, True, "FFFFFF"),
                       Alignment(horizontal=h, vertical="center", indent=0 if h == "center" else 1))

    def item(self, n, it, precio_global=None):
        nl = lineas(it["descripcion"], DESC_CHARS)
        r = self.fila(H_ITEM + H_ITEM_LINE * (nl - 1))
        self.unir(f"B{r}:D{r}")
        sub = it.get("incluye") or []
        self.celda(f"A{r}", n, font(10, False, GREY), Alignment(horizontal="center", vertical="center"))
        self.celda(f"B{r}", it["descripcion"], font(10, bool(sub)),
                   Alignment(horizontal="left", vertical="center", indent=1, wrap_text=nl > 1))
        self.celda(f"E{r}", it.get("modelo") or "", font(9, False, GREY), Alignment(horizontal="left", vertical="center", indent=1))
        self.celda(f"F{r}", it.get("cantidad"), font(10), Alignment(horizontal="center", vertical="center"))
        al = Alignment(horizontal="right", vertical="center", indent=1)
        if precio_global is not None:
            self.celda(f"H{r}", precio_global, font(10, True), al, MONEY)
        else:
            self.celda(f"G{r}", it.get("precio_unitario"), font(10), al, MONEY)
            self.celda(f"H{r}", f'=IF(F{r}*G{r}=0,"",F{r}*G{r})', font(10, True), al, MONEY)
        self.celda(f"G{r}", None, None, None, MONEY)
        ultima = r
        if sub:
            ri = self.fila(H_SUB)
            self.unir(f"B{ri}:D{ri}")
            self.celda(f"B{ri}", "Incluye:", font(9, True, GREY), Alignment(horizontal="left", vertical="center", indent=1))
            for comp in sub:
                rc = self.fila(H_SUB)
                self.unir(f"B{rc}:D{rc}")
                self.celda(f"B{rc}", "•  " + comp, font(9, False, GREY), Alignment(horizontal="left", vertical="center", indent=2))
            ultima = self.fila(4)
        for col in "ABCDEFGH":
            self.ws[f"{col}{ultima}"].border = Border(bottom=HAIR)
        return r, ultima

    def vacia(self):
        r = self.fila(H_EMPTY)
        self.unir(f"B{r}:D{r}")
        for col in "ABCDEFGH":
            self.ws[f"{col}{r}"].border = Border(bottom=HAIR)
        self.celda(f"G{r}", None, font(10), Alignment(horizontal="right", vertical="center", indent=1), MONEY)
        self.celda(f"H{r}", f'=IF(F{r}*G{r}=0,"",F{r}*G{r})', font(10, True),
                   Alignment(horizontal="right", vertical="center", indent=1), MONEY)
        return r

    def continua(self):
        r = self.fila(22)
        self.unir(f"A{r}:H{r}")
        self.celda(f"A{r}", "Continúa en la hoja siguiente  →", font(8, False, GREY, italic=True),
                   Alignment(horizontal="right", vertical="center", indent=1))

    def cierre(self, primera, ultima):
        d = self.d
        iva = d.get("iva", 0)
        self.fila(12)
        rs = self.fila(18)
        ri = self.fila(18)
        self.fila(5)
        rt = self.fila(28)
        # condiciones a la izquierda de los totales
        for a, b, lab, val in (("A", "B", "CONDICIÓN DE PAGO", d.get("condicion_pago") or "100% contado"),
                               ("C", "E", "COTIZACIÓN USD BNA", money(d["cotizacion_usd"]))):
            self.unir(f"{a}{rs}:{b}{rs}")
            self.unir(f"{a}{ri}:{b}{ri}")
            self.celda(f"{a}{rs}", lab, font(7, True, GREY), Alignment(horizontal="left", vertical="bottom", indent=1))
            self.celda(f"{a}{ri}", val, font(10, True), Alignment(horizontal="left", vertical="center", indent=1))
        iva_txt = {0: "IVA  (0%)", 21: "IVA  (21%)", 10.5: "IVA  (10,5%)"}[iva]
        for r, lab, formula in ((rs, "Subtotal", f"=SUM(H{primera}:H{ultima})"),
                                (ri, iva_txt, 0 if iva == 0 else f"=H{rs}*{iva / 100}")):
            self.unir(f"F{r}:G{r}")
            self.celda(f"F{r}", lab, font(9, False, GREY), Alignment(horizontal="right", vertical="center", indent=1))
            self.celda(f"H{r}", formula, font(10), Alignment(horizontal="right", vertical="center", indent=1), MONEY)
            for col in "FGH":
                self.ws[f"{col}{r}"].border = Border(bottom=HAIR)
        self.unir(f"F{rt}:G{rt}")
        self.pintar(rt, "FGH", SLATE)
        self.celda(f"F{rt}", "TOTAL", font(11, True, "FFFFFF"), Alignment(horizontal="right", vertical="center", indent=1))
        self.celda(f"H{rt}", f"=H{rs}+H{ri}", font(14, True, "FFFFFF"),
                   Alignment(horizontal="right", vertical="center", indent=1), MONEY)
        obs = (d.get("observaciones") or "").strip()
        if obs:
            self.fila(14)
            rc = self.fila(14)
            self.unir(f"A{rc}:H{rc}")
            self.celda(f"A{rc}", "OBSERVACIONES", font(7, True, GREY), Alignment(horizontal="left", vertical="bottom", indent=1))
            ro = self.fila(max(30, 16 + 13 * lineas(obs, OBS_CHARS)))
            self.unir(f"A{ro}:H{ro}")
            self.celda(f"A{ro}", obs, font(10), Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True))
            self.pintar(ro, "ABCDEFGH", TINT)

    def pie(self):
        ws = self.ws
        ws.oddFooter.left.text = "REFRIVAN  ·  Documento no válido como factura."
        ws.oddFooter.right.text = "Hoja &P de &N"
        for parte in (ws.oddFooter.left, ws.oddFooter.right):
            parte.size = 7
            parte.font = "Arial,Regular"
            parte.color = GREY


def alto_cierre(d):
    h = 12 + 18 + 18 + 5 + 28
    obs = (d.get("observaciones") or "").strip()
    if obs:
        h += 14 + 14 + max(30, 16 + 13 * lineas(obs, OBS_CHARS))
    return h


H_ENC1 = 4 + 20 + 30 + 18 + 16 + 6 + 5 + 14
H_CLIENTE = 6 + 13 + 20 + 4 + 13 + 20 + 6 + 14
H_ENC2 = 4 + 17 + 17 + 8 + 4 + 12
H_TITULO, H_CONTINUA = 22, 22


def paginar(items, d, cap):
    """Reparte los ítems en hojas. Devuelve una lista de listas de índices."""
    cierre = alto_cierre(d)
    hojas, i, n = [], 0, len(items)
    while True:
        disp = cap - (H_ENC1 + H_CLIENTE if not hojas else H_ENC2) - H_TITULO
        resto = sum(alto_item(items[k]) for k in range(i, n))
        if resto + cierre <= disp:
            hojas.append(list(range(i, n)))
            return hojas
        tomados, usado = [], 0
        while i < n and usado + alto_item(items[i]) + H_CONTINUA <= disp:
            usado += alto_item(items[i])
            tomados.append(i)
            i += 1
        if i == n and tomados:  # entran los ítems pero no el cierre: el último pasa a la hoja siguiente
            tomados.pop()
            i -= 1
        if not tomados:
            if i >= n:
                raise SystemExit("El cierre no entra ni en una hoja nueva")
            tomados.append(i)  # un ítem más alto que la hoja: va solo
            i += 1
        hojas.append(tomados)


def armar(d, cap):
    items = d["items"]
    hojas = paginar(items, d, cap)
    h = Hoja(d)
    h.encabezado()
    h.cliente()
    primera = ultima = None
    n = 0
    for k, idx in enumerate(hojas):
        if k:
            h.encabezado_continuacion(k + 1, len(hojas))
        h.titulo_tabla()
        for i in idx:
            n += 1
            pg = d.get("precio_global") if n == 1 else None
            r0, r1 = h.item(n, items[i], pg)
            primera = primera or r0
            ultima = r1
        if k < len(hojas) - 1:
            h.continua()
            h.ws.row_breaks.append(Break(id=h.r - 1))
    if len(hojas) == 1:  # una sola hoja: completar la tabla con renglones vacíos si hay lugar
        libre = cap - H_ENC1 - H_CLIENTE - H_TITULO - sum(alto_item(it) for it in items) - alto_cierre(d)
        vacias = min(max(0, 7 - len(items)), int(libre // H_EMPTY))
        for _ in range(vacias):
            ultima = h.vacia()
    h.cierre(primera, ultima)
    h.pie()
    ws = h.ws
    ws.print_area = f"A1:H{h.r - 1}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.45, bottom=0.8, header=0.2, footer=0.3)
    return h.wb, len(hojas)


def escala_medida(pdf):
    """Escala con la que LibreOffice imprimió la hoja: distancia real entre CLIENTE y DIRECCIÓN / nominal."""
    import pymupdf
    ys = {w[4]: w[1] for w in pymupdf.open(pdf)[0].get_text("words") if w[4] in ("CLIENTE", "DIRECCIÓN")}
    if len(ys) < 2:
        return None
    return (ys["DIRECCIÓN"] - ys["CLIENTE"]) / (13 + 20 + 4)


def a_pdf(xlsx, carpeta):
    perfil = tempfile.mkdtemp()
    subprocess.run(["soffice", f"-env:UserInstallation=file://{perfil}", "--headless", "--convert-to", "pdf",
                    "--outdir", carpeta, xlsx], check=True, capture_output=True, timeout=180)
    shutil.rmtree(perfil, ignore_errors=True)
    return os.path.splitext(xlsx)[0] + ".pdf"


def paginas(pdf):
    import pymupdf
    return len(pymupdf.open(pdf))


def main():
    datos_path, carpeta = sys.argv[1], sys.argv[2]
    png = "--png" in sys.argv
    with open(datos_path, encoding="utf-8") as f:
        d = json.load(f)
    d.setdefault("fecha", datetime.date.today().strftime("%d/%m/%Y"))
    d.setdefault("cliente", {})
    if d.get("iva") not in (0, 21, 10.5):
        raise SystemExit("iva tiene que ser 0, 21 o 10.5")
    os.makedirs(carpeta, exist_ok=True)
    nombre = os.path.join(carpeta, f"Presupuesto N{d['numero']}.xlsx")
    # 1) primera pasada para medir la escala real de impresión; 2) armar con el alto útil real
    cap = PAGE_PT
    wb, _ = armar(d, cap)
    wb.save(nombre)
    esc = escala_medida(a_pdf(nombre, carpeta))
    if esc:
        cap = PAGE_PT / esc - 12
        print(f"Escala de impresión: {esc:.2f}")
    for _ in range(6):
        wb, esperadas = armar(d, cap)
        wb.save(nombre)
        pdf = a_pdf(nombre, carpeta)
        reales = paginas(pdf)
        if reales == esperadas:
            break
        cap -= 25  # LibreOffice cortó antes de lo calculado: achicar el alto útil y rearmar
    else:
        raise SystemExit(f"No pude cuadrar las hojas ({reales} en vez de {esperadas})")
    if png:
        import pymupdf
        for i, p in enumerate(pymupdf.open(pdf)):
            p.get_pixmap(dpi=110).save(os.path.join(carpeta, f"Presupuesto N{d['numero']}-hoja{i + 1}.png"))
    print(f"OK  {nombre}\nOK  {pdf}\nHojas: {reales}")


if __name__ == "__main__":
    main()
