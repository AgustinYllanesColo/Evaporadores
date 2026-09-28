"""Genera la plantilla nueva PRESUPUESTO REFRIVAN.xlsx (mismo mapa de celdas que PRESUPUESTO LIMPIO)."""
import sys
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

LOGO = sys.argv[1] if len(sys.argv) > 1 else "logo.png"
OUT = sys.argv[2] if len(sys.argv) > 2 else "PRESUPUESTO REFRIVAN.xlsx"

SLATE, RED, ORANGE, TEAL = "283137", "BF222D", "EB9E06", "19988A"
GREY, LINE, TINT = "5B6670", "D5DBE0", "F1F4F6"
MONEY = '[$-2C0A]"$ "#,##0'

def font(size=10, bold=False, color=SLATE, italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)

def fill(c):
    return PatternFill("solid", fgColor=c)

hair = Side(style="thin", color=LINE)
strong = Side(style="medium", color=SLATE)

wb = Workbook()
ws = wb.active
ws.title = "Presupuesto"
ws.sheet_view.showGridLines = False

for col, w in {"A": 8, "B": 30, "C": 9, "D": 9, "E": 12, "F": 7, "G": 14, "H": 16}.items():
    ws.column_dimensions[col].width = w

heights = {1: 6, 2: 20, 3: 30, 4: 18, 5: 6, 6: 13, 7: 13, 8: 4, 9: 12,
           10: 24, 11: 24, 12: 24, 13: 6, 14: 6, 15: 24,
           25: 8, 26: 20, 27: 20, 28: 28, 29: 12, 30: 18, 31: 24, 32: 10,
           33: 18, 34: 64, 35: 10, 36: 14, 37: 14}
for r in range(16, 25):
    heights[r] = 30
for r, h in heights.items():
    ws.row_dimensions[r].height = h

# --- Encabezado: logo a la izquierda, número y fecha a la derecha
logo = XLImage(LOGO)
ratio = logo.width / logo.height
logo.height = 62
logo.width = round(62 * ratio)
ws.add_image(logo, "A2")

ws.merge_cells("G2:H2")
ws["G2"] = "PRESUPUESTO"
ws["G2"].font = font(10, True, GREY)
ws["G2"].alignment = Alignment(horizontal="right", vertical="bottom")

ws["G3"] = "N°"
ws["G3"].font = font(10, False, GREY)
ws["G3"].alignment = Alignment(horizontal="right", vertical="center")
ws["H3"] = 0
ws["H3"].font = font(20, True)
ws["H3"].number_format = "0"
ws["H3"].alignment = Alignment(horizontal="right", vertical="center")

ws["G4"] = "Fecha"
ws["G4"].font = font(9, False, GREY)
ws["G4"].alignment = Alignment(horizontal="right", vertical="center")
ws["H4"] = None
ws["H4"].font = font(10, True)
ws["H4"].number_format = "dd/mm/yyyy"
ws["H4"].alignment = Alignment(horizontal="right", vertical="center")

for r, txt in {6: "Marco Avellaneda 960, Remedios de Escalada, Buenos Aires  ·  www.refrivan.com.ar",
               7: "Tel: +54 011 3840-0340  ·  refrivansitioweb@gmail.com"}.items():
    ws.merge_cells(f"A{r}:H{r}")
    ws[f"A{r}"] = txt
    ws[f"A{r}"].font = font(8, False, GREY)
    ws[f"A{r}"].alignment = Alignment(horizontal="left", vertical="center")

# franja tricolor de la marca
for cols, c in (("ABC", RED), ("DEF", ORANGE), ("GH", TEAL)):
    for col in cols:
        ws[f"{col}8"].fill = fill(c)

# --- Datos del cliente (recuadro gris claro)
for r in (10, 11, 12):
    for col in "ABCDEFGH":
        ws[f"{col}{r}"].fill = fill(TINT)
    ws.merge_cells(f"B{r}:D{r}")
    ws.merge_cells(f"E{r}:F{r}")
    ws.merge_cells(f"G{r}:H{r}")
for r, (l1, l2) in {10: ("Cliente", "Contacto"), 11: ("Dirección", "Localidad"), 12: ("Teléfono", "Provincia")}.items():
    for cell, txt in ((f"A{r}", l1), (f"E{r}", l2)):
        ws[cell] = txt
        ws[cell].font = font(9, False, GREY)
        ws[cell].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for cell in (f"B{r}", f"G{r}"):
        ws[cell].font = font(10, r == 10 and cell.startswith("B"))
        ws[cell].alignment = Alignment(horizontal="left", vertical="center")

# --- Tabla de ítems
ws.merge_cells("B15:D15")
heads = {"A15": ("N°", "center"), "B15": ("DESCRIPCIÓN", "left"), "E15": ("MODELO", "left"),
         "F15": ("CANT.", "center"), "G15": ("PRECIO UNIT.", "right"), "H15": ("SUBTOTAL", "right")}
for col in "ABCDEFGH":
    ws[f"{col}15"].fill = fill(SLATE)
for cell, (txt, h) in heads.items():
    ws[cell] = txt
    ws[cell].font = font(8, True, "FFFFFF")
    ws[cell].alignment = Alignment(horizontal=h, vertical="center", indent=1 if h != "center" else 0)

for r in range(16, 25):
    ws.merge_cells(f"B{r}:D{r}")
    for col in "ABCDEFGH":
        ws[f"{col}{r}"].border = Border(bottom=hair)
        ws[f"{col}{r}"].font = font(10)
    ws[f"A{r}"].font = font(10, False, GREY)
    ws[f"A{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"B{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws[f"E{r}"].font = font(9, False, GREY)
    ws[f"E{r}"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws[f"F{r}"].alignment = Alignment(horizontal="center", vertical="center")
    ws[f"G{r}"].number_format = MONEY
    ws[f"G{r}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws[f"H{r}"] = f'=IF(F{r}*G{r}=0,"",F{r}*G{r})'
    ws[f"H{r}"].number_format = MONEY
    ws[f"H{r}"].font = font(10, True)
    ws[f"H{r}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)

# --- Totales
for r, label in {26: "Subtotal", 27: "IVA  (0%)", 28: "TOTAL"}.items():
    ws.merge_cells(f"F{r}:G{r}")
    ws[f"F{r}"] = label
    ws[f"H{r}"].number_format = MONEY
    ws[f"F{r}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    ws[f"H{r}"].alignment = Alignment(horizontal="right", vertical="center", indent=1)
    if r < 28:
        ws[f"F{r}"].font = font(10, False, GREY)
        ws[f"H{r}"].font = font(10)
        for col in "FGH":
            ws[f"{col}{r}"].border = Border(bottom=hair)
ws["H26"] = "=SUM(H16:H24)"
ws["H27"] = 0
ws["H28"] = "=H26+H27"
ws["F28"].font = font(11, True)
ws["H28"].font = font(14, True)
for col in "FGH":
    ws[f"{col}28"].fill = fill(TINT)
    ws[f"{col}28"].border = Border(top=strong, bottom=strong)

# --- Condiciones
ws.merge_cells("A30:H30")
ws["A30"] = "CONDICIONES"
ws["A30"].font = font(8, True, GREY)
ws["A30"].alignment = Alignment(vertical="bottom")
ws.merge_cells("A31:C31")
ws.merge_cells("E31:G31")
for col in "ABCDEFGH":
    ws[f"{col}31"].fill = fill(TINT)
ws["A31"] = "Cotización USD BNA"
ws["E31"] = "Condición de pago"
for cell in ("A31", "E31"):
    ws[cell].font = font(9, False, GREY)
    ws[cell].alignment = Alignment(horizontal="right" if cell == "E31" else "left", vertical="center", indent=1)
ws["D31"] = "$ 0"
ws["H31"] = "100% contado"
for cell in ("D31", "H31"):
    ws[cell].font = font(10, True)
ws["D31"].alignment = Alignment(horizontal="left", vertical="center")
ws["H31"].alignment = Alignment(horizontal="right", vertical="center", indent=1)

# --- Observaciones
ws.merge_cells("A33:H33")
ws["A33"] = "OBSERVACIONES"
ws["A33"].font = font(8, True, GREY)
ws["A33"].alignment = Alignment(vertical="bottom")
ws.merge_cells("A34:H34")
ws["A34"].font = font(10)
ws["A34"].alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
for col in "ABCDEFGH":
    ws[f"{col}34"].border = Border(top=hair, bottom=hair,
                                   left=hair if col == "A" else None,
                                   right=hair if col == "H" else None)

# --- Pie
ws.merge_cells("A36:H36")
ws["A36"] = "Documento no válido como factura."
ws["A36"].font = font(8, False, GREY, italic=True)
ws["A36"].alignment = Alignment(horizontal="center", vertical="center")
for col in "ABCDEFGH":
    ws[f"{col}36"].border = Border(top=hair)
ws.merge_cells("A37:H37")
ws["A37"] = "REFRIVAN  ·  Home · Gastro · Bazar"
ws["A37"].font = font(8, True, GREY)
ws["A37"].alignment = Alignment(horizontal="center", vertical="center")

# --- Impresión: A4 vertical, una página de ancho, centrado
ws.print_area = "A1:H37"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True
ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.55, bottom=0.5, header=0, footer=0)

wb.save(OUT)
print("ok", OUT)
